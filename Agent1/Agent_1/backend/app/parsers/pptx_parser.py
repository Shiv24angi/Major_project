from pptx import Presentation


def extract_pptx_text(file_path: str) -> dict:
    presentation = Presentation(file_path)

    slides = []
    full_text_parts = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        texts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    texts.append(text)

        combined_text = "\n".join(texts)

        slides.append(
            {
                "slide_number": slide_number,
                "text": combined_text
            }
        )

        if combined_text:
            full_text_parts.append(
                f"--- SLIDE {slide_number} ---\n"
                f"{combined_text}"
            )

    return {
        "slide_count": len(slides),
        "slides": slides,
        "full_text": "\n\n".join(full_text_parts)
    }